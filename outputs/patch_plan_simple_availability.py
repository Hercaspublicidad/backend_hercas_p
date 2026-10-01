from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import openpyxl

root = Path(r"D:\api_0doo\outputs")
source = root / "plan-oficial-backup-20260930-simple-availability.xlsx"
target = root / "plan-oficial-staged-20260930-simple-availability.xlsx"
edits = {
    "xl/worksheets/sheet2.xml": {
        "G20": "Parcial: consulta pública simple en desarrollo",
        "H20": "30/09: 85 vallas habilitadas provisionalmente en hercas-platform-dev por autorización de Carlos. RPC pública devuelve código, nombre y estado diario disponible/reservada/ocupada para 46 días inclusivos; vista /disponibilidad conectada y verificada en navegador. Las 85 figuran pendientes de revisión por Katherine; cero ocupaciones activas y ninguna fecha verificada. No usar para confirmar contratación aún.",
        "G50": "Parcial: consulta de disponibilidad conectada",
        "H50": "30/09: /disponibilidad del frontend integraci0n consulta Supabase, permite buscar entre 85 vallas y elegir fechas; muestra estados diarios y refresca cada 30 segundos. Verificado con 14 días disponibles en APT-003. Vista solo de lectura. Pendiente revisión de Katherine y aceptación de datos reales.",
    },
    "xl/worksheets/sheet3.xml": {
        "D11": "30/09: website y calendario Supabase son centro de consulta de disponibilidad. Tres estados diarios: disponible, reservada y ocupada. Por autorización de Carlos se habilitaron provisionalmente las 85 vallas en hercas-platform-dev para probar /disponibilidad; las 85 siguen marcadas para revisión por Katherine, sin last_verified_at ni ocupaciones activas. Vista frontend conectada y comprobada en navegador; refresca cada 30 segundos. Antes de confirmar una contratación hay que conciliar contratos, bloqueos y fechas reales. Reservas, pagos e integración comercial quedan para etapa posterior. Evidencia: migraciones 20260930223345 y 20260930223637, docs/disponibilidad-centro-control.md.",
    },
}
ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with ZipFile(source, "r") as zin, ZipFile(target, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in edits:
            tree = etree.fromstring(data)
            for address, value in edits[item.filename].items():
                cells = tree.xpath(f'.//m:c[@r="{address}"]', namespaces=ns)
                if len(cells) != 1:
                    raise RuntimeError(f"{item.filename} {address}: {len(cells)} matches")
                cell = cells[0]
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, "{%s}is" % ns["m"])
                text = etree.SubElement(inline, "{%s}t" % ns["m"])
                text.text = value
            data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
        zout.writestr(item, data)

before = openpyxl.load_workbook(source)
after = openpyxl.load_workbook(target)
for sheet in before.sheetnames:
    a, b = before[sheet], after[sheet]
    for row in a:
        for cell in row:
            new = b[cell.coordinate]
            if cell.data_type == "f" and (new.data_type != "f" or new.value != cell.value):
                raise RuntimeError(f"Formula changed: {sheet}!{cell.coordinate}")
for path, values in edits.items():
    sheet = "Requerimientos" if "sheet2" in path else "Decisiones"
    for address, expected in values.items():
        if after[sheet][address].value != expected:
            raise RuntimeError(f"Value mismatch: {sheet}!{address}")
print(f"Staged {sum(map(len, edits.values()))} cells; all formulas preserved: {target}")
