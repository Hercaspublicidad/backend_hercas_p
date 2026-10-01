from openpyxl import load_workbook

PATH = r"D:\api_0doo\outputs\supabase-permissions-20260925\Hercas - Plan de trabajo Anderson y tu.before.xlsx"
workbook = load_workbook(PATH, data_only=False)
for name in ["Requerimientos", "Decisiones", "Plan a dos", "Cierre de etapas"]:
    sheet = workbook[name]
    print("---", name)
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True), start=1):
        values = [str(value)[:140] if value is not None else "" for value in row]
        if any(token in " ".join(values) for token in ["BE-005", "INT-003", "SER-007", "FE-014", "D-016", "Etapa 1", "Auth", "autentic"]):
            print(index, " | ".join(values))
