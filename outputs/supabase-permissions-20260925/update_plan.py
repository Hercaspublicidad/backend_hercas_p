from copy import copy
from openpyxl import load_workbook

INPUT = r"D:\api_0doo\outputs\supabase-permissions-20260925\Hercas - Plan de trabajo Anderson y tu.before.xlsx"
OUTPUT = r"D:\api_0doo\outputs\supabase-permissions-20260925\Hercas - Plan de trabajo Anderson y tu.updated.xlsx"
workbook = load_workbook(INPUT)
requirements = workbook["Requerimientos"]

# Preserve the workbook structure and edit only the affected evidence/status cells.
updates = {
    11: {
        7: "Parcial: migración remota aplicada y validada",
        8: "25/09: RPC admin_update_role_permissions aplicado en Supabase; anónimo bloqueado, autenticado habilitado y trigger de auditoría verificado. UI permite editar módulos por rol. Falta validar recorrido real completo de correo/invitación.",
        16: "Pendiente: validar envío real y aceptar invitación; evidencia SQL y prueba local de autorización disponibles.",
    },
    45: {
        7: "Parcial: panel administrativo implementado",
        8: "25/09: el panel permite editar los módulos de cada rol mediante RPC autorizado; mantiene users.manage para systems_admin. El envío de invitaciones informa claramente si el backend no tiene la clave privada configurada.",
        16: "Pendiente: prueba de invitación por correo real con backend configurado.",
    },
    62: {
        7: "Parcial: Auth y URLs configuradas; entrega pendiente",
        8: "25/09: la invitación puede prepararse, pero el backend no tiene SUPABASE_SECRET_KEY configurada para llamar a Auth y enviar el correo. Configurarla solo en el servidor; no exponerla en Next.js.",
        16: "Bloqueado para envío real hasta configurar la clave privada de Supabase en el backend y ejecutar prueba de entrega/aceptación.",
    },
    82: {
        7: "Parcial: control de permisos aplicado; hardening pendiente",
        8: "25/09: Supabase confirma RPC administrativo no ejecutable por anon y auditoría de cambios de permisos. Advisor detecta protección de contraseñas filtradas desactivada; activar antes de cierre de etapa.",
        16: "Pendiente: activar leaked password protection, MFA privilegiado, WAF y límites distribuidos con evidencia de despliegue.",
    },
}
for row, cells in updates.items():
    for column, value in cells.items():
        requirements.cell(row=row, column=column).value = value

decisions = workbook["Decisiones"]
decisions.cell(row=22, column=9).value = "En revisión"
decisions.cell(row=22, column=16).value = (
    "25/09: migración administrativa y panel de módulos aplicados. Falta configurar la clave privada de Supabase en backend para entrega real de correos; no se aprueba cierre sin esa evidencia."
)

workbook.save(OUTPUT)
print(OUTPUT)
