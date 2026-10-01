from openpyxl import load_workbook

PATH = r"D:\api_0doo\outputs\supabase-permissions-20260925\Hercas - Plan de trabajo Anderson y tu.updated.xlsx"
workbook = load_workbook(PATH, data_only=False)
requirements = workbook["Requerimientos"]
for row in (11, 45, 62, 82):
    print(row, [requirements.cell(row=row, column=column).value for column in (1, 7, 8, 16)])
print("D-016", [workbook["Decisiones"].cell(row=22, column=column).value for column in (1, 9, 16, 17)])
print("formulas", workbook["Decisiones"].cell(row=22, column=17).value, workbook["Plan a dos"].cell(row=8, column=8).value)
for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "#REF!" in cell.value:
                raise RuntimeError(f"Unexpected #REF! in {sheet.title}!{cell.coordinate}")
print("verified")
