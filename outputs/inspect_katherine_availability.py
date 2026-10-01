from collections import Counter
from datetime import datetime
from pathlib import Path
import json
import openpyxl

path = Path(r"D:\Disponible Semana del 27 de Septiembre al 03 de Octubre -2026- precios actulizados 2026 - Katherine.xls")
with path.open("rb") as stream:
    workbook = openpyxl.load_workbook(stream, data_only=True)
sheet = workbook.worksheets[1]
records = []
for index in range(3, sheet.max_row + 1):
    row = [sheet.cell(index, column).value for column in range(1, 25)]
    code = row[1]
    if not isinstance(code, str) or not code.strip():
        continue
    record = {
        "row": index, "code": code.strip().upper(), "state": str(row[6] or "").strip(),
        "price_month": row[7], "price_2026": row[11], "price_2026_alt": row[20],
        "customer": row[12], "campaign": row[13], "uninstall": row[14],
        "start": row[15], "end": row[16], "advisor": row[17], "note": row[18],
    }
    for key in ("uninstall", "start", "end"):
        if isinstance(record[key], datetime): record[key] = record[key].date().isoformat()
    records.append(record)
print("rows", len(records), "states", dict(Counter(record["state"] for record in records)))
print("dated", sum(bool(record["start"] and record["end"]) for record in records))
print("missing_dates_occupied", [record["code"] for record in records if "ARREND" in record["state"].upper() and not (record["start"] and record["end"])])
print("duplicates", {code: count for code, count in Counter(record["code"] for record in records).items() if count > 1})
print("future_or_current", json.dumps([record for record in records if record["end"] and record["end"] >= "2026-10-01"], ensure_ascii=False, default=str))
Path(r"D:\api_0doo\outputs\katherine-availability-20261001.json").write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
