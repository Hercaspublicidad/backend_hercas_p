from openpyxl import load_workbook

INPUT = r"D:\api_0doo\outputs\fase-1-auth-20260925\Hercas - Plan de trabajo Anderson y tu.before.xlsx"
OUTPUT = r"D:\api_0doo\outputs\fase-1-auth-20260925\Hercas - Plan de trabajo Anderson y tu.updated.xlsx"
workbook = load_workbook(INPUT)
requirements = workbook["Requerimientos"]
requirements.cell(row=82, column=7).value = "Parcial: controles de Auth aplicados; hardening pendiente"
requirements.cell(row=82, column=8).value = (
    "25/09: registro público desactivado, confirmación de email activa y Supabase configurado para exigir contraseña actual al cambiarla. "
    "RPC administrativo bloqueado para anon y auditado. La protección contra contraseñas filtradas requiere plan Pro; evaluar actualización o control compensatorio."
)
requirements.cell(row=82, column=16).value = (
    "Pendiente: MFA para administradores, WAF y límites distribuidos. Protección contra contraseñas filtradas bloqueada por plan Free."
)
workbook.save(OUTPUT)
print(OUTPUT)
