from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import openpyxl

root = Path(r"D:\api_0doo\outputs")
source = root / "plan-oficial-backup-20261001-bulk-workbench.xlsx"
target = root / "plan-oficial-staged-20261001-bulk-workbench.xlsx"
edits = {
    "xl/worksheets/sheet2.xml": {
        "H20": "01/10: lectura Supabase de 85 vallas en matriz pública compacta de hasta 31 días, tres estados y reporte de vencimientos del mes sin cliente. El website /disponibilidad muestra la tabla y refresca cada 30 segundos; consulta verificada en navegador. La disponibilidad inicial de las 85 vallas sigue pendiente de revisión por Katherine; hoy 84 disponibles y 1 ocupada en APT-005, sin cliente registrado. No confirmar contratación con datos aún sin conciliar.",
        "H21": "01/10: Supabase availability_bulk_workbench aplicado. Comercial/gestor puede seleccionar hasta 100 vallas y asignar disponible, reservada manual persistente u ocupada durante rango de hasta 12 meses; cliente y motivo requeridos, soporte para ocupada. Controla permisos, propiedad, revisión y cruce; el cambio es atómico y conserva tramos fuera del rango. Website integraci0n ofrece tabla compacta, selección masiva, atajos 1/3/6/12 meses y cliente por valla para personal. Prueba transaccional revertida: reservada 32 días, ocupada 8, liberados 3 intermedios; estados correctos. Falta conciliación real por Katherine.",
        "H50": "01/10: /disponibilidad rediseñada: 85 vallas en tabla de fechas pequeñas, filtros, edición en el mismo panel, estados disponible/reservada/ocupada, selección masiva hasta 12 meses y vencimientos mensuales. El equipo autenticado ve cliente por valla cuando está registrado; la vista pública no ve nombres. TypeScript, ESLint y Next build pasaron; navegador verificó tabla, selección de dos vallas, atajo 3m y reporte. BD de desarrollo aplicada y flujo SQL reversible verificado. Continúa pendiente revisión de ocupaciones reales por Katherine y aceptación integral del negocio.",
    },
    "xl/worksheets/sheet3.xml": {
        "D11": "01/10: website y Supabase son centro de disponibilidad. Se aplicó availability_bulk_workbench: matriz diaria de 85 vallas, tres estados, reservas manuales persistentes, edición masiva atómica hasta 100 vallas/12 meses, cliente visible sólo al personal autorizado y vencimientos mensuales sin cliente para público. /disponibilidad en integraci0n pasó a tabla compacta con edición integrada y atajos mensuales. Prueba SQL revertida confirmó conservar tramos al reservar, ocupar y liberar; build frontend y navegador pasaron. Existe 1 ocupación activa APT-005 sin cliente registrado; las 85 vallas siguen marcadas pendientes de revisión de Katherine. No confirmar contratación hasta conciliar datos reales. Evidencia: docs/disponibilidad-centro-control.md.",
    },
}

ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with ZipFile(source, "r") as zin, ZipFile(target, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in edits:
            tree = etree.fromstring(data)
            for address, value in edits[item.filename].items():
                cells = tree.xpath(f'.//m:c[@r="{address}"]', namespaces=ns)
                if len(cells) != 1:
                    raise RuntimeError(f"{item.filename} {address}: {len(cells)} matches")
                cell = cells[0]
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, "{%s}is" % ns["m"])
                text = etree.SubElement(inline, "{%s}t" % ns["m"])
                text.text = value
            data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
        zout.writestr(item, data)

before = openpyxl.load_workbook(source)
after = openpyxl.load_workbook(target)
for sheet in before.sheetnames:
    a, b = before[sheet], after[sheet]
    for row in a:
        for cell in row:
            new = b[cell.coordinate]
            if cell.data_type == "f" and (new.data_type != "f" or new.value != cell.value):
                raise RuntimeError(f"Formula changed: {sheet}!{cell.coordinate}")
for path, values in edits.items():
    sheet = "Requerimientos" if "sheet2" in path else "Decisiones"
    for address, expected in values.items():
        if after[sheet][address].value != expected:
            raise RuntimeError(f"Value mismatch: {sheet}!{address}")
print(f"Staged {sum(map(len, edits.values()))} cells; formulas preserved: {target}")
