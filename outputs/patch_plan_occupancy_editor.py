from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import openpyxl

root = Path(r"D:\api_0doo\outputs")
source = root / "plan-oficial-backup-20261001-occupancy-editor.xlsx"
target = root / "plan-oficial-staged-20261001-occupancy-editor.xlsx"
edits = {
    "xl/worksheets/sheet2.xml": {
        "G21": "Parcial: ocupación comercial desde website",
        "H21": "01/10: /disponibilidad en frontend integraci0n permite a comercial o gestor autenticado marcar como ocupada una valla en fechas elegidas, con motivo y referencia de soporte, y liberar con motivo. La base impide cruces y valida permisos, propietario y revisión. Prueba reversible en APT-003: 14 días pasaron de disponible a ocupada y volvieron a disponible; 0 ocupaciones activas, historial cancelado conservado. Falta asignar/verificar cuenta de Katherine y conciliar datos reales.",
        "G50": "Parcial: consulta y control de ocupación conectados",
        "H50": "01/10: /disponibilidad muestra las 85 vallas y estados diarios. A cuentas sales/availability_manager/systems_admin presenta gestión: ocupar intervalo con motivo y soporte, listar ocupaciones y liberar con motivo; en prueba de navegador los estados cambiaron y se restauraron. La consulta pública sigue marcada provisional hasta revisión del inventario por Katherine. Pago, reserva cliente y aceptación integral pendientes.",
    },
    "xl/worksheets/sheet3.xml": {
        "D11": "01/10: website y calendario Supabase son centro de disponibilidad. La vista /disponibilidad permite consultar tres estados diarios y a comerciales/gestor ocupar o liberar fechas mediante RPC auditadas. Prueba reversible APT-003: disponible -> ocupada -> disponible, 0 ocupaciones activas; historial cancelado. Las 85 vallas fueron habilitadas provisionalmente por Carlos y siguen pendientes de revisión por Katherine, sin last_verified_at. No se identificó cuenta activa de Katherine por nombre; debe recibir availability_manager para liberar ocupaciones ajenas. No confirmar contratación antes de conciliar contratos y bloqueos reales. Pago y reserva cliente quedan posteriores. Evidencia: docs/disponibilidad-centro-control.md.",
    },
}
ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with ZipFile(source, "r") as zin, ZipFile(target, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in edits:
            tree = etree.fromstring(data)
            for address, value in edits[item.filename].items():
                cells = tree.xpath(f'.//m:c[@r="{address}"]', namespaces=ns)
                if len(cells) != 1:
                    raise RuntimeError(f"{item.filename} {address}: {len(cells)} matches")
                cell = cells[0]
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, "{%s}is" % ns["m"])
                text = etree.SubElement(inline, "{%s}t" % ns["m"])
                text.text = value
            data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
        zout.writestr(item, data)

before = openpyxl.load_workbook(source)
after = openpyxl.load_workbook(target)
for sheet in before.sheetnames:
    a, b = before[sheet], after[sheet]
    for row in a:
        for cell in row:
            new = b[cell.coordinate]
            if cell.data_type == "f" and (new.data_type != "f" or new.value != cell.value):
                raise RuntimeError(f"Formula changed: {sheet}!{cell.coordinate}")
for path, values in edits.items():
    sheet = "Requerimientos" if "sheet2" in path else "Decisiones"
    for address, expected in values.items():
        if after[sheet][address].value != expected:
            raise RuntimeError(f"Value mismatch: {sheet}!{address}")
print(f"Staged {sum(map(len, edits.values()))} cells; all formulas preserved: {target}")
