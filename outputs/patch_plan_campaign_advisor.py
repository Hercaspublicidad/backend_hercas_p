from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import openpyxl

root = Path(r"D:\api_0doo\outputs")
source = root / "plan-oficial-backup-20261001-weekly-import.xlsx"
target = root / "plan-oficial-staged-20261001-weekly-import.xlsx"
source_workbook = openpyxl.load_workbook(source)
append = {
    "xl/worksheets/sheet2.xml": {
        "H20": " 01/10 carga semanal: Excel 27/09-03/10 de Katherine incorporado (112 filas, 85 coincidencias). 23 periodos con fechas y cuatro vallas no ofertables; APT-005 liberada por decisión de Carlos. Se quitaron 85 marcas de revisión general. Precios históricos sólo internos; precio público a cargo de Gabriel, tráfico sin verificar. Tres filas disponibles con cliente hasta 05/10 se conservan ocupadas hasta ese fin para evitar doble oferta.",
        "H21": " 01/10 carga semanal: 23 ocupaciones con cliente y fechas; campaña y asesor histórico preservados como fuente. Corrección parcial probada en transacción revertida sin perder el asesor. SPT-001 arrendada sin fechas queda no ofertable; asesores del Excel requieren vinculación con cuentas Auth para nuevas operaciones. Sin filas de prueba persistidas.",
        "H50": " 01/10 carga semanal: pantalla sin aviso general de Katherine; vencimientos de octubre (9), ocupación de 23% para 01-14/10 calculada sobre 81 vallas ofertables. Dos XLSX descargados y abiertos: disponibilidad con 58 vallas libres todo el rango, 85 filas internas y 23 asignaciones; vencimientos con 9 filas. Build, TypeScript, ESLint y navegador verificados. Precio público pendiente de Gabriel; tráfico pendiente de fuente.",
    },
    "xl/worksheets/sheet3.xml": {
        "D11": " 01/10 carga semanal de Katherine: 112 filas preservadas como fuente privada con hash; 85 códigos conciliados, 23 intervalos fechados, cuatro activos no ofertables. APT-005 liberada por decisión de Carlos; tres filas disponibles con fechas de cliente se mantienen ocupadas hasta 05/10. Revisión general inicial retirada. Indicadores excluyen cuatro bloqueadas; dos Excel validados. Precio público a cargo de Gabriel y tráfico sin verificación. Asesores importados por nombre hasta vincular Auth. Evidencia: docs/disponibilidad-centro-control.md. Aceptación integral parcial.",
    },
}
ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with ZipFile(source, "r") as zin, ZipFile(target, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in append:
            tree = etree.fromstring(data)
            for address, addition in append[item.filename].items():
                cells = tree.xpath(f'.//m:c[@r="{address}"]', namespaces=ns)
                if len(cells) != 1:
                    raise RuntimeError(f"Missing {item.filename}!{address}")
                cell = cells[0]
                prior = source_workbook["Requerimientos" if "sheet2" in item.filename else "Decisiones"][address].value
                if not isinstance(prior, str):
                    raise RuntimeError(f"Unexpected {address}")
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, "{%s}is" % ns["m"])
                value = etree.SubElement(inline, "{%s}t" % ns["m"])
                value.text = prior + addition
            data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
        zout.writestr(item, data)

after = openpyxl.load_workbook(target)
for sheet in source_workbook.sheetnames:
    before_sheet, after_sheet = source_workbook[sheet], after[sheet]
    for row in before_sheet:
        for cell in row:
            new = after_sheet[cell.coordinate]
            if cell.data_type == "f" and (new.data_type != "f" or new.value != cell.value):
                raise RuntimeError(f"Formula changed: {sheet}!{cell.coordinate}")
for path, changes in append.items():
    sheet = "Requerimientos" if "sheet2" in path else "Decisiones"
    for address, addition in changes.items():
        assert after[sheet][address].value == source_workbook[sheet][address].value + addition
print(target)
